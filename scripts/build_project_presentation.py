"""Build the Vietnamese project overview deck.

Install the optional authoring dependency before running:

    python -m pip install python-pptx
    python scripts/build_project_presentation.py

The generated PPTX intentionally uses only editable PowerPoint shapes and text.
"""

from __future__ import annotations

import argparse
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


BG = RGBColor(10, 18, 32)
SURFACE = RGBColor(22, 33, 53)
SURFACE_ALT = RGBColor(31, 46, 70)
TEXT = RGBColor(244, 247, 251)
MUTED = RGBColor(170, 184, 204)
TEAL = RGBColor(45, 205, 178)
BLUE = RGBColor(91, 151, 255)
CORAL = RGBColor(255, 119, 112)
YELLOW = RGBColor(248, 202, 93)
GREEN = RGBColor(101, 211, 145)
FONT = "Aptos"
MONO = "Aptos Mono"


def rgb(value: RGBColor) -> str:
    return str(value)


class Deck:
    def __init__(self) -> None:
        self.prs = Presentation()
        self.prs.slide_width = Inches(13.333)
        self.prs.slide_height = Inches(7.5)
        self.page = 0

    def blank(self, *, footer: bool = True):
        slide = self.prs.slides.add_slide(self.prs.slide_layouts[6])
        self.page += 1
        bg = slide.background.fill
        bg.solid()
        bg.fore_color.rgb = BG
        if footer:
            self.text(
                slide,
                0.55,
                7.12,
                11.7,
                0.2,
                "Meeting Task Pipeline  •  17.08.2026",
                8.5,
                MUTED,
            )
            self.text(
                slide,
                12.25,
                7.08,
                0.55,
                0.22,
                str(self.page),
                9,
                MUTED,
                align=PP_ALIGN.RIGHT,
            )
        return slide

    def shape(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        *,
        fill: RGBColor = SURFACE,
        line: RGBColor | None = None,
        radius: bool = True,
    ):
        kind = MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE
        box = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
        box.fill.solid()
        box.fill.fore_color.rgb = fill
        box.line.color.rgb = line or fill
        return box

    def text(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        value: str,
        size: float,
        color: RGBColor = TEXT,
        *,
        bold: bool = False,
        align=PP_ALIGN.LEFT,
        font: str = FONT,
        valign=MSO_ANCHOR.TOP,
        margin: float = 0.02,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(margin)
        frame.margin_right = Inches(margin)
        frame.margin_top = Inches(margin)
        frame.margin_bottom = Inches(margin)
        frame.vertical_anchor = valign
        paragraph = frame.paragraphs[0]
        paragraph.text = value
        paragraph.alignment = align
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color
        return box

    def bullets(
        self,
        slide,
        x: float,
        y: float,
        w: float,
        h: float,
        items: list[str],
        *,
        size: float = 17,
        color: RGBColor = TEXT,
        accent: RGBColor = TEAL,
    ):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.margin_left = Inches(0.08)
        frame.margin_right = Inches(0.04)
        frame.margin_top = Inches(0.03)
        for index, item in enumerate(items):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = item
            paragraph.font.name = FONT
            paragraph.font.size = Pt(size)
            paragraph.font.color.rgb = color
            paragraph.space_after = Pt(10)
            paragraph.level = 0
            paragraph.text = "•  " + paragraph.text
            if paragraph.runs:
                paragraph.runs[0].font.color.rgb = color
        return box

    def header(self, slide, title: str, kicker: str | None = None) -> None:
        self.shape(slide, 0.55, 0.42, 0.08, 0.62, fill=TEAL, radius=False)
        if kicker:
            self.text(slide, 0.82, 0.38, 11.6, 0.22, kicker.upper(), 9, TEAL, bold=True)
            self.text(slide, 0.8, 0.64, 11.9, 0.52, title, 26, TEXT, bold=True)
        else:
            self.text(slide, 0.8, 0.47, 11.9, 0.62, title, 28, TEXT, bold=True)

    def title_slide(self) -> None:
        slide = self.blank(footer=False)
        self.shape(slide, 0, 0, 0.18, 7.5, fill=TEAL, radius=False)
        self.text(slide, 0.78, 0.76, 4.6, 0.3, "PROJECT TECHNICAL OVERVIEW", 11, TEAL, bold=True)
        self.text(slide, 0.76, 1.32, 11.5, 1.55, "Meeting Task\nPipeline", 43, TEXT, bold=True)
        self.text(
            slide,
            0.8,
            3.25,
            8.3,
            1.0,
            "Từ transcript tiếng Việt/Anh đến task ledger có evidence —\nPython-first, AI-last, tích hợp Power Automate.",
            21,
            MUTED,
        )
        self.shape(slide, 9.55, 1.05, 2.75, 3.95, fill=SURFACE_ALT, line=BLUE)
        self.text(slide, 9.92, 1.48, 2.05, 0.28, "DEFAULT PATH", 9, BLUE, bold=True, align=PP_ALIGN.CENTER)
        self.text(slide, 9.86, 2.05, 2.16, 0.5, "V1", 32, TEXT, bold=True, align=PP_ALIGN.CENTER)
        self.text(slide, 9.86, 2.7, 2.16, 0.35, "Context: assist", 14, MUTED, align=PP_ALIGN.CENTER)
        self.text(slide, 9.86, 3.18, 2.16, 0.35, "86 cases", 14, MUTED, align=PP_ALIGN.CENTER)
        self.text(slide, 9.86, 3.66, 2.16, 0.35, "FastAPI + PA", 14, MUTED, align=PP_ALIGN.CENTER)
        self.text(slide, 0.8, 6.72, 8.0, 0.35, "Technical briefing • Architecture • Evaluation • Operations", 11, MUTED)

    def section(self, number: str, title: str, subtitle: str) -> None:
        slide = self.blank()
        self.text(slide, 0.78, 1.08, 2.0, 1.2, number, 68, TEAL, bold=True)
        self.text(slide, 2.92, 1.22, 9.4, 0.88, title, 34, TEXT, bold=True)
        self.text(slide, 2.94, 2.28, 8.7, 1.0, subtitle, 19, MUTED)
        self.shape(slide, 2.94, 3.62, 7.9, 0.06, fill=TEAL, radius=False)

    def two_columns(
        self,
        title: str,
        left_title: str,
        left_items: list[str],
        right_title: str,
        right_items: list[str],
        *,
        kicker: str | None = None,
        left_color: RGBColor = BLUE,
        right_color: RGBColor = TEAL,
        note: str | None = None,
    ) -> None:
        slide = self.blank()
        self.header(slide, title, kicker)
        for x, label, items, color in (
            (0.72, left_title, left_items, left_color),
            (6.78, right_title, right_items, right_color),
        ):
            self.shape(slide, x, 1.42, 5.82, 4.92, fill=SURFACE, line=color)
            self.text(slide, x + 0.28, 1.72, 5.2, 0.42, label, 18, color, bold=True)
            self.bullets(slide, x + 0.25, 2.27, 5.22, 3.75, items, size=15.5)
        if note:
            self.text(slide, 0.82, 6.56, 11.7, 0.32, note, 10.5, YELLOW, bold=True, align=PP_ALIGN.CENTER)

    def flow(self, title: str, nodes: list[tuple[str, str]], *, kicker: str | None = None) -> None:
        slide = self.blank()
        self.header(slide, title, kicker)
        rows = [nodes[:5], nodes[5:]] if len(nodes) > 5 else [nodes]
        for row_index, row in enumerate(rows):
            y = 1.72 + row_index * 2.35
            box_w = 2.18 if len(row) >= 5 else min(2.6, 11.4 / len(row))
            gap = 0.25
            total = len(row) * box_w + (len(row) - 1) * gap
            start = (13.333 - total) / 2
            for index, (label, detail) in enumerate(row):
                x = start + index * (box_w + gap)
                color = [BLUE, TEAL, YELLOW, CORAL, GREEN][index % 5]
                self.shape(slide, x, y, box_w, 1.52, fill=SURFACE, line=color)
                self.text(slide, x + 0.16, y + 0.23, box_w - 0.32, 0.4, label, 15, color, bold=True, align=PP_ALIGN.CENTER)
                self.text(slide, x + 0.16, y + 0.72, box_w - 0.32, 0.54, detail, 10.5, MUTED, align=PP_ALIGN.CENTER)
                if index < len(row) - 1:
                    self.text(slide, x + box_w + 0.015, y + 0.58, gap - 0.03, 0.35, "→", 17, MUTED, bold=True, align=PP_ALIGN.CENTER)
        if len(rows) == 2:
            self.text(slide, 6.05, 3.36, 1.2, 0.36, "↓", 22, MUTED, bold=True, align=PP_ALIGN.CENTER)

    def metric_cards(self, title: str, cards: list[tuple[str, str, str, RGBColor]], note: str) -> None:
        slide = self.blank()
        self.header(slide, title, "Evidence-backed status")
        width = 2.84
        for index, (value, label, detail, color) in enumerate(cards):
            x = 0.68 + index * 3.13
            self.shape(slide, x, 1.6, width, 3.42, fill=SURFACE, line=color)
            self.text(slide, x + 0.18, 1.98, width - 0.36, 0.72, value, 29, color, bold=True, align=PP_ALIGN.CENTER)
            self.text(slide, x + 0.18, 2.82, width - 0.36, 0.55, label, 15, TEXT, bold=True, align=PP_ALIGN.CENTER)
            self.text(slide, x + 0.26, 3.55, width - 0.52, 0.9, detail, 11, MUTED, align=PP_ALIGN.CENTER)
        self.shape(slide, 1.1, 5.43, 11.1, 0.88, fill=SURFACE_ALT, line=YELLOW)
        self.text(slide, 1.36, 5.69, 10.6, 0.36, note, 13, YELLOW, bold=True, align=PP_ALIGN.CENTER)

    def table(
        self,
        title: str,
        headers: list[str],
        rows: list[list[str]],
        widths: list[float],
        *,
        kicker: str | None = None,
        font_size: float = 11.5,
    ) -> None:
        slide = self.blank()
        self.header(slide, title, kicker)
        x0, y0 = 0.67, 1.45
        row_h = min(0.72, 5.28 / (len(rows) + 1))
        x = x0
        for header, width in zip(headers, widths):
            self.shape(slide, x, y0, width, row_h, fill=SURFACE_ALT, line=BG, radius=False)
            self.text(slide, x + 0.08, y0 + 0.1, width - 0.16, row_h - 0.18, header, 11, TEAL, bold=True, valign=MSO_ANCHOR.MIDDLE)
            x += width
        for row_index, row in enumerate(rows):
            x = x0
            y = y0 + (row_index + 1) * row_h
            fill = SURFACE if row_index % 2 == 0 else RGBColor(26, 39, 61)
            for value, width in zip(row, widths):
                self.shape(slide, x, y, width, row_h, fill=fill, line=BG, radius=False)
                self.text(slide, x + 0.08, y + 0.08, width - 0.16, row_h - 0.14, value, font_size, TEXT, valign=MSO_ANCHOR.MIDDLE)
                x += width


def build_deck(output: Path) -> None:
    deck = Deck()
    deck.title_slide()

    deck.two_columns(
        "Tóm tắt điều hành",
        "Hệ thống đang làm tốt",
        [
            "Contract và pipeline có ranh giới rõ: Python quyết định final state; AI chỉ sinh event trung gian.",
            "Evidence, diagnostics và unresolved windows giúp audit thay vì che mơ hồ.",
            "API đồng bộ, async job, Docker và Azure Functions dùng cùng FastAPI app.",
            "Task Ledger đã tách identity, aliases, owner/deadline và terminal state.",
        ],
        "Điều chưa được hiểu sai",
        [
            "Chất lượng canonical chưa đạt production; pass-case vẫn thấp.",
            "Gate C xác nhận transport/provider/replay, không chứng minh toàn bộ extraction đã đúng.",
            "Job store chỉ ở memory; quick tunnel chỉ dành cho demo/test.",
            "Power Automate cần upsert, review queue và metadata ngày/title chuẩn.",
        ],
        kicker="Executive summary",
        note="Khuyến nghị: demo có kiểm soát + tiếp tục quality work; chưa tự động ghi task production không qua review.",
    )

    deck.section("01", "Bài toán và ranh giới", "Pipeline nhận transcript, suy ra task còn hiệu lực và giữ đủ dấu vết để người dùng kiểm tra.")

    deck.flow(
        "Từ đầu vào đến đầu ra công khai",
        [
            ("INPUT", "TXT / VTT / SRT\nMeeting metadata + note"),
            ("PARSE", "Caption → turn → sentence → clause"),
            ("EXTRACT", "Rule events + bounded AI events"),
            ("REDUCE", "Global Task Ledger theo chronology"),
            ("OUTPUT", "Summary + tasks + evidence + diagnostics"),
        ],
        kicker="Public contract",
    )

    deck.two_columns(
        "Input và output contract",
        "MeetingInput",
        [
            "meeting_id, meeting_title, meeting_date",
            "transcript_raw, file_name, speaker_aliases",
            "meeting_note là optional; source: SECRETARY / PARTICIPANT / MANUAL / AUTO_OVERVIEW",
            "Ngày họp phải được gửi khi test date logic.",
        ],
        "PipelineResult",
        [
            "meeting_title và summary template-based",
            "tasks: task_name, assignee, start/due date, raw due text, evidence, status",
            "diagnostics: rule/AI/window/batch/fallback/replay counts",
            "unresolved_window_ids là tín hiệu review bắt buộc.",
        ],
        kicker="Stable API",
    )

    deck.section("02", "Kiến trúc runtime", "V1 là production path; V2 vẫn là shadow/experimental và cung cấp một số helper dùng chung.")

    deck.table(
        "Bản đồ codebase",
        ["Khu vực", "Vai trò", "Điểm cần nhớ"],
        [
            ["backend/app", "FastAPI + toàn bộ runtime", "main.py chọn provider; pipeline.py orchestration"],
            ["ingestion → annotation", "Parse, normalize, clause, cue/date", "Giữ order_index và stable IDs"],
            ["candidate + ai", "Window, routing, rule/AI extractors", "AI chỉ xử lý candidate mơ hồ"],
            ["reduction + dates", "Link event, Task Ledger, resolve date", "Reducer global chạy tuần tự"],
            ["output", "Task, evidence, summary, diagnostics", "Giữ public contract tương thích"],
            ["data/validation", "86 case ground truth", "Source of truth sau code/tests"],
            ["evaluation + scripts", "Report, trace, evaluator, runbook", "Report local ≠ live provider"],
            ["power-automate + sp365", "Flow, schema, Lists mapping", "Orchestration, không duplicate business logic"],
        ],
        [1.8, 4.15, 5.38],
        kicker="Repository topology",
        font_size=10.5,
    )

    deck.flow(
        "Pipeline V1 — chuỗi xử lý chính",
        [
            ("INGEST", "Parser + caption dedup"),
            ("PREPROCESS", "Speaker/turn/sentence/clause"),
            ("ANNOTATE", "Cue flags + date mentions"),
            ("CONTEXT", "Topic + note grounding"),
            ("CANDIDATE", "Window + route"),
            ("EVENTS", "Rule + AI + recap"),
            ("AUTHORITY", "Filter + event dedup"),
            ("LEDGER", "Link + reduce + terminal guard"),
            ("DATES", "Deterministic resolver"),
            ("SERIALIZE", "Tasks + evidence + diagnostics"),
        ],
        kicker="Chronological and auditable",
    )

    deck.two_columns(
        "Candidate routing: RULE, AI hay CONTEXT?",
        "Local rule",
        [
            "Assignment/commitment rõ, action cụ thể, score đủ cao.",
            "Mutation chỉ local khi target ổn định, thường có task label.",
            "Recap/snapshot có authority riêng khi đủ scope và cấu trúc.",
            "Brainstorm, question, progress, hypothetical thường không create.",
        ],
        "AI fallback",
        [
            "Ambiguous reference, mutation thiếu target, note cue mơ hồ.",
            "Chỉ nhận bounded clauses + task memory tối đa 5 candidates.",
            "anchor_clause_id phải thuộc primary clauses.",
            "Provider lỗi → giữ rule result + đánh dấu unresolved; meeting không fail.",
        ],
        kicker="AI-last",
    )

    deck.two_columns(
        "Meeting Context và Meeting Note policy",
        "Context compaction",
        [
            "Clause được xếp MANDATORY / LIKELY / CONTEXT / NOISE.",
            "Classification chỉ chọn context gửi AI; không prune rule input.",
            "Focus + safety clauses luôn giữ; budget 56 là soft limit.",
            "Window xa nhau không gộp để tránh coreference xuyên đoạn.",
        ],
        "Meeting Note",
        [
            "Human positive note có thể tạo event khi assertion + action + owner/task guard.",
            "AUTO_OVERVIEW chỉ context/summary, không local create.",
            "Note thiếu row không xóa task; transcript mutation vẫn có authority.",
            "AI event luôn phải anchor về transcript clause.",
        ],
        kicker="Safety and grounding",
        right_color=YELLOW,
    )

    deck.section("03", "Task Ledger và final state", "Event không phải task cuối; reducer tích lũy mutation theo chronology rồi mới serialize active snapshot.")

    deck.flow(
        "Event contract và reducer",
        [
            ("CREATE", "TASK_CREATE / COMMITMENT"),
            ("OWNER", "ASSIGN / REASSIGN"),
            ("DEADLINE", "SET / REPLACE"),
            ("TERMINAL", "CANCEL / REJECT"),
            ("SNAPSHOT", "Chỉ task ACTIVE được xuất"),
        ],
        kicker="Sequential state machine",
    )

    deck.two_columns(
        "Task identity và reconciliation",
        "Identity layer",
        [
            "Canonical task_id ổn định trong ledger.",
            "Alias/label/reference được tích lũy, không dùng owner/date làm identity chính.",
            "TaskLinker chấm candidate và có confidence/margin guard.",
            "Ambiguous mutation không tự gắn bừa vào task gần nhất.",
        ],
        "Mutation layer",
        [
            "Owner và deadline là state có thể ghi đè.",
            "Terminal state chặn replay tạo lại cùng identity.",
            "Recap có thể reconcile state nhưng phải qua authority/scope.",
            "Final serializer chỉ đọc active ledger snapshot.",
        ],
        kicker="Task Ledger refactor",
    )

    deck.table(
        "Start-date và deadline semantics deterministic",
        ["Cụm thời gian", "Kết quả", "Nguyên tắc"],
        [
            ["task bắt đầu từ 20/02", "start = 20/02", "Task start explicit có priority cao nhất"],
            ["Ngày họp: 17/01/2026", "meeting anchor = 17/01", "Request → transcript → note → processing date"],
            ["trước thứ Sáu", "Thứ Năm", "before D không có giờ → D - 1 ngày"],
            ["trước 18h thứ Sáu", "Thứ Sáu", "Có giờ/cuối ngày → vẫn là D"],
            ["vào / by / chậm nhất D", "D", "Inclusive deadline"],
            ["15/01 không có năm", "Năm họp hoặc rollover", "Nếu đã qua meeting date → năm sau"],
            ["sau 3 ngày", "start + 3 calendar days", "Calendar duration có thể resolve"],
            ["3 ngày làm việc", "due_date rỗng", "Không tự suy diễn lịch làm việc"],
            ["sau khi khách duyệt", "due_date rỗng", "Event-dependent; giữ raw text"],
        ],
        [2.75, 2.65, 5.93],
        kicker="No calendar guessing",
        font_size=11,
    )

    deck.section("04", "API, hosting và vận hành", "Một FastAPI application được host bằng Uvicorn, Docker hoặc Azure Functions; async job là main PA flow.")

    deck.table(
        "API surface",
        ["Method", "Endpoint", "Use case"],
        [
            ["GET", "/health", "Health probe"],
            ["POST", "/api/v1/transcripts/preprocess", "Inspect parser/clauses"],
            ["POST", "/api/v1/meetings/process", "Sync JSON"],
            ["POST", "/api/v1/meetings/process-file", "Sync file"],
            ["POST", "/api/v1/meetings/jobs/process", "Async JSON submit"],
            ["POST", "/api/v1/meetings/jobs/process-file", "Async binary — main PA flow"],
            ["GET", "/api/v1/meetings/jobs/{job_id}", "Poll queued/running/succeeded/failed"],
        ],
        [1.0, 5.15, 5.18],
        kicker="FastAPI contract",
        font_size=11.5,
    )

    deck.flow(
        "Async job lifecycle",
        [
            ("SUBMIT", "202 + job_id + status_url"),
            ("QUEUE", "In-memory job store"),
            ("RUN", "V1 + selected provider"),
            ("POLL", "queued / running"),
            ("RESULT", "succeeded/result hoặc failed/error"),
        ],
        kicker="Idempotent within one process",
    )

    deck.two_columns(
        "Hosting và provider selection",
        "Hosting",
        [
            "Local: Uvicorn backend.app.main:app.",
            "Docker: cùng app trên port 8000.",
            "Azure Functions: func.AsgiFunctionApp qua compatibility export.",
            "Quick Tunnel: demo/test public endpoint; URL không durable.",
        ],
        "Provider priority",
        [
            "1. OPENAI_API_KEY → Responses API.",
            "2. Azure AI Foundry chat endpoint.",
            "3. Generic AI_FALLBACK_ENDPOINT.",
            "4. DisabledAiClient → rule-only + unresolved diagnostics.",
        ],
        kicker="One application, multiple hosts",
    )

    deck.section("05", "Power Automate integration", "Power Automate điều phối file, job và Microsoft Lists; business logic vẫn ở backend.")

    deck.flow(
        "Power Automate main flow",
        [
            ("TRIGGER", "SharePoint / OneDrive file"),
            ("SUBMIT", "POST binary + metadata headers"),
            ("POLL", "GET status_url với backoff"),
            ("UPSERT", "MI Meetings + Task Proposals"),
            ("REVIEW", "NeedsReview / approval / publish"),
        ],
        kicker="Orchestration, not extraction",
    )

    deck.table(
        "File request headers cho Power Automate",
        ["Header", "Trạng thái", "Giá trị"],
        [
            ["Content-Type", "Required", "application/octet-stream"],
            ["X-API-Key", "Required khi backend enforce", "Secret từ environment variable / Key Vault"],
            ["X-File-Name-Base64", "Required", "base64(Name) — hỗ trợ Unicode an toàn"],
            ["Embedded metadata", "Preferred", "ID/title/date nằm trong cùng file package"],
            ["X-Meeting-Id", "Optional override", "Dùng cho raw client cũ"],
            ["X-Meeting-Title/Date", "Optional override", "Header > package > context > processing date"],
        ],
        [2.8, 2.0, 6.53],
        kicker="Unicode-safe metadata",
        font_size=11.2,
    )

    deck.two_columns(
        "Lists mapping và idempotency",
        "MI Meetings",
        [
            "MeetingId là unique business key.",
            "Upsert Processing → Writing → Success/Failed.",
            "Lưu summary, diagnostics, unresolved count, pipeline/model.",
            "NeedsReview nếu unresolved > 0 hoặc policy yêu cầu.",
        ],
        "MI Task Proposals",
        [
            "ProposalKey = MeetingId|TaskSequence là unique key.",
            "Upsert task proposal; không Create blindly khi retry.",
            "Lưu evidence, dates, raw due text, assignee, review status.",
            "Chỉ publish downstream sau human review ở giai đoạn PoC.",
        ],
        kicker="Storage-level deduplication",
    )

    deck.two_columns(
        "Power Automate: các guard bắt buộc",
        "Đã harden trong repo",
        [
            "Job key bao gồm content + title/date/id/file/aliases.",
            "JSON schemas cho submit/status envelope.",
            "Base64 headers tránh lỗi Unicode của HTTP connector.",
            "Docs thống nhất status_url, secure inputs/outputs và upsert.",
        ],
        "Còn phải làm ở tenant",
        [
            "Tạo unique columns/index và Configure run after.",
            "Bật Secure Inputs/Outputs cho submit, poll và parse result.",
            "Xử lý 404 sau backend restart: resubmit với cùng metadata.",
            "Thay quick tunnel bằng Azure HTTPS endpoint trước production.",
        ],
        kicker="Repo readiness vs tenant configuration",
        right_color=CORAL,
    )

    deck.section("06", "Quality, testing và evaluation", "Ground truth reviewed + deterministic comparator vẫn là chuẩn; live provider được đánh giá như một biến thể có metadata đầy đủ.")

    deck.metric_cards(
        "Canonical quality snapshot",
        [
            ("15 / 86", "V1 with notes", "P 0.4040\nR 0.5776\nField 0.8333", BLUE),
            ("16 / 86", "V1 without notes", "P 0.3793\nR 0.5162\nField 0.8473", TEAL),
            ("0.4091", "Gate C replay P", "R 0.5848\nField 0.8333", YELLOW),
            ("234", "Unexpected tasks", "Gate C replay\n173 unresolved", CORAL),
        ],
        "Không diễn giải gate transport/provider/replay là quality production gate; false positives và identity vẫn là trọng tâm.",
    )

    deck.table(
        "Corpus và test strategy",
        ["Lớp", "Phạm vi", "Mục tiêu"],
        [
            ["Unit", "Parser, annotation, dates, rules, linker, reducer", "Rule positive luôn có negative test"],
            ["Integration", "Pipeline, note/context, jobs, contracts", "Giữ boundary giữa các stage"],
            ["End-to-end", "FastAPI sync/async/file/auth/idempotency", "Bảo vệ public contract"],
            ["Targeted cases", "Case lỗi hoặc semantic cluster", "Chứng minh fix đúng lỗi cụ thể"],
            ["Wave W1–W5", "86 reviewed cases", "Phát hiện regression theo độ dài/feature"],
            ["Full A/B", "With/without notes, local/live", "Đo quality và tác động Meeting Note"],
        ],
        [2.0, 4.35, 5.0],
        kicker="Gates from cheap to expensive",
        font_size=11,
    )

    deck.two_columns(
        "Có nên làm sản phẩm đánh giá?",
        "Giữ comparator canonical",
        [
            "Blocking gate phải deterministic và reproducible.",
            "Precision/recall/field metrics không phụ thuộc mood của model.",
            "Ground truth chỉ đổi sau review transcript/final state.",
            "Mọi run ghi rõ version, note mode, provider, model và prompt.",
        ],
        "Xây Evaluation Workbench nhỏ",
        [
            "Expected/actual cạnh nhau, transcript + evidence + ledger.",
            "Semantic matching chỉ gợi ý alignment, không tự chấm pass.",
            "Human adjudication + error taxonomy + audit patch.",
            "Dashboard theo wave/label/length/note/provider.",
        ],
        kicker="Recommendation",
        note="Không chỉ “so sánh”, nhưng cũng không xây AI judge thay ground truth.",
    )

    deck.flow(
        "Evaluation Workbench — workflow đề xuất",
        [
            ("COMPARE", "Deterministic expected ↔ actual"),
            ("SUGGEST", "Semantic alignment non-blocking"),
            ("INSPECT", "Transcript + evidence + events"),
            ("ADJUDICATE", "Human decision + taxonomy"),
            ("EXPORT", "Versioned audit/patch + issue"),
        ],
        kicker="Measure correctly, learn faster",
    )

    deck.section("07", "Roadmap và quyết định", "Ưu tiên quality, durable operations và review workflow trước khi tự động hóa production.")

    deck.table(
        "Roadmap thực tế",
        ["Ưu tiên", "Hạng mục", "Exit condition"],
        [
            ["P0", "Power Automate demo flow + unique keys + secure actions", "Public submit/poll/upsert smoke pass"],
            ["P0", "False-create / task identity / long-distance mutation", "Targeted + wave + full A/B không regression"],
            ["P1", "Evaluation Workbench read-only + taxonomy", "Reviewer xử lý case nhanh, audit được"],
            ["P1", "Durable job store/queue + TTL + multi-instance", "Restart không mất job; idempotency durable"],
            ["P1", "Canonical live-AI A/B", "Run metadata + usage + replay artifact đầy đủ"],
            ["P2", "Blind/dev corpus thực tế", "Không tune trực tiếp trên canonical test"],
            ["P2", "V2 parity hoặc thu hẹp dependency", "Quyết định promote, retain shadow hoặc retire"],
        ],
        [1.05, 6.25, 4.05],
        kicker="From PoC to production",
        font_size=10.8,
    )

    deck.two_columns(
        "Demo runbook",
        "Backend",
        [
            "Set POWER_AUTOMATE_API_KEY, PIPELINE_VERSION=v1, CONTEXT=assist.",
            "Start: uvicorn backend.app.main:app --port 8010.",
            "Health: GET /health.",
            "Nếu không có OPENAI_API_KEY: demo là rule-only và có unresolved diagnostics.",
        ],
        "Power Automate",
        [
            "Set API base URL từ HTTPS tunnel/Azure endpoint.",
            "Submit self-contained package + Base64 file-name header.",
            "Poll status_url đến terminal status.",
            "Upsert Lists; route NeedsReview; kiểm tra evidence trước publish.",
        ],
        kicker="Operational checklist",
    )

    deck.two_columns(
        "Known limitations",
        "Runtime / platform",
        [
            "Job store in-memory, chưa TTL/persistence/multi-instance.",
            "Quick tunnel URL thay đổi và không có SLA.",
            "V2 chưa parity V1; V1 còn dùng helper dưới v2.",
            "C# sample client chưa theo async job + Meeting Note.",
        ],
        "Quality / data",
        [
            "Pass rate canonical thấp; unexpected tasks còn nhiều.",
            "Chưa có blind/dev corpus thực tế.",
            "Current canonical reports chủ yếu deterministic; live-AI cần run chuẩn mới.",
            "Trace chứa raw transcript/evidence, phải coi là dữ liệu nhạy cảm.",
        ],
        kicker="What prevents production sign-off",
        left_color=CORAL,
        right_color=YELLOW,
    )

    slide = deck.blank()
    deck.header(slide, "Quyết định đề xuất", "Close")
    deck.shape(slide, 0.82, 1.58, 11.7, 1.12, fill=SURFACE_ALT, line=TEAL)
    deck.text(
        slide,
        1.18,
        1.88,
        10.95,
        0.5,
        "Demo Power Automate được — production automation chưa được.",
        24,
        TEAL,
        bold=True,
        align=PP_ALIGN.CENTER,
    )
    deck.bullets(
        slide,
        1.28,
        3.22,
        10.8,
        2.45,
        [
            "Giữ V1 làm main path và comparator deterministic làm blocking gate.",
            "Dùng Evaluation Workbench để review nhanh, đúng và có audit trail.",
            "Hoàn tất Power Automate tenant guards, sau đó demo end-to-end có human review.",
            "Chỉ mở production write-back khi quality gate và durable job infrastructure đạt yêu cầu.",
        ],
        size=18,
    )
    deck.text(slide, 1.18, 6.28, 10.95, 0.36, "Nguồn: code/tests, data/validation, evaluation artifacts và docs hiện hành.", 11, MUTED, align=PP_ALIGN.CENTER)

    output.parent.mkdir(parents=True, exist_ok=True)
    deck.prs.save(output)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--output",
        type=Path,
        default=Path("docs/presentation/meeting-intelligent-project-overview.pptx"),
    )
    args = parser.parse_args()
    build_deck(args.output)
    print(f"Wrote {args.output}")


if __name__ == "__main__":
    main()
